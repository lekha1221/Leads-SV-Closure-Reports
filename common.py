"""Shared Athena + Excel helpers for leads_sv_closure_reports notebooks."""
from __future__ import annotations

from datetime import date, timedelta
from pathlib import Path
from typing import Any

import awswrangler as wr
import boto3
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows

ROOT = Path(__file__).resolve().parent
ENV_PATHS = [ROOT.parent / ".env", Path(".env")]

ANALYSIS_START = date(2026, 1, 1)
ANALYSIS_END = date.today()
START_TS = f"TIMESTAMP '{ANALYSIS_START.isoformat()} 00:00:00'"
END_TS = f"TIMESTAMP '{(ANALYSIS_END + timedelta(days=1)).isoformat()} 00:00:00'"
CLOSURE_STAGE = "Sales Closure"


def load_env() -> dict[str, str]:
    for p in ENV_PATHS:
        if not p.exists():
            continue
        out: dict[str, str] = {}
        for line in p.read_text().splitlines():
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            k, v = line.split("=", 1)
            out[k.strip()] = v.strip()
        print("Loaded", p)
        return out
    raise FileNotFoundError("No .env found")


def make_aws_session(env: dict[str, str] | None = None) -> boto3.Session:
    """SSO profile session — run `aws sso login --profile <name>` when expired."""
    env = env or load_env()
    profile = (env.get("AWS_PROFILE") or "asbl").strip()
    region = (env.get("AWS_REGION") or "ap-south-1").strip()
    session = boto3.Session(profile_name=profile, region_name=region)
    try:
        session.client("sts").get_caller_identity()
    except Exception as e:
        raise RuntimeError(
            f"AWS SSO login required. Run:\n  aws sso login --profile {profile}\n({e})"
        ) from e
    return session


def athena_session() -> tuple[Any, str, str, str]:
    env = load_env()
    session = make_aws_session(env)
    return (
        session,
        env.get("ATHENA_DATABASE", "leadsdb-prod"),
        env.get("ATHENA_WORKGROUP", "primary"),
        env.get("ATHENA_OUTPUT_LOCATION", "s3://metabase-asbl/athena-query-results/"),
    )


_SESSION = None
_DB = _WG = _S3 = ""


def q(sql: str, database: str | None = None) -> pd.DataFrame:
    global _SESSION, _DB, _WG, _S3
    if _SESSION is None:
        _SESSION, _DB, _WG, _S3 = athena_session()
    return wr.athena.read_sql_query(
        sql=sql,
        database=database or _DB,
        boto3_session=_SESSION,
        workgroup=_WG,
        s3_output=_S3,
        ctas_approach=False,
    )


def opt(clause: str, values: Any) -> str:
    """Optional filter (Metabase `[[ ... ]]` equivalent): '' when no values given.

    `%s` in `clause` is replaced by a quoted, comma-separated SQL list.
    """
    if not values:
        return ""
    if isinstance(values, str):
        values = [values]
    quoted = ", ".join("'" + str(v).replace("'", "''") + "'" for v in values)
    return clause.replace("%s", quoted)


def write_excel(path: Path, sheets: dict[str, pd.DataFrame]) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    wb = Workbook()
    wb.remove(wb.active)
    for name, df in sheets.items():
        ws = wb.create_sheet(name[:31])
        for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), start=1):
            for c_idx, val in enumerate(row, start=1):
                cell = ws.cell(r_idx, c_idx, None if pd.isna(val) else val)
                cell.border = thin
                if r_idx == 1:
                    cell.fill = header_fill
                    cell.font = header_font
        for col in ws.columns:
            letter = get_column_letter(col[0].column)
            width = min(40, max(10, max(len(str(c.value or "")) for c in col) + 2))
            ws.column_dimensions[letter].width = width
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    print("Wrote", path, "→", list(sheets))
